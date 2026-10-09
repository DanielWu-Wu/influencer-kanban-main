from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
SOURCE = Path(r'C:\Users\Admin\AppData\Local\Temp')
OUT = ROOT / 'Apify YouTube 字幕 API 获取教程.docx'
FONT = r'C:\Windows\Fonts\msyh.ttc'


def annotated(name, boxes):
    image = Image.open(SOURCE / name).convert('RGB')
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(FONT, 34)
    for number, box in enumerate(boxes, 1):
        draw.rounded_rectangle(box, radius=14, outline='#e11937', width=7)
        x, y = box[:2]
        draw.ellipse((x-16, y-18, x+30, y+28), fill='#e11937')
        draw.text((x+7, y+5), str(number), anchor='mm', font=font, fill='white')
    target = SOURCE / ('marked-' + name)
    image.save(target, quality=95)
    return target


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.7)
section.bottom_margin = Cm(1.7)
section.left_margin = Cm(2)
section.right_margin = Cm(2)
styles = doc.styles
styles['Normal'].font.name = 'Microsoft YaHei'
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(7)
styles['Title'].font.name = 'Microsoft YaHei'
styles['Title'].font.size = Pt(24)
styles['Title'].font.color.rgb = RGBColor(0, 0, 0)
for key in ['Heading 1', 'Heading 2']:
    styles[key].font.name = 'Microsoft YaHei'
    styles[key].font.color.rgb = RGBColor(28, 55, 89)


def line(text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    return p


def fig(path, caption):
    doc.add_picture(str(path), width=Cm(16.8))
    p = doc.add_paragraph(caption)
    p.style = 'Caption'


doc.add_paragraph('Apify YouTube 字幕 API 获取教程', style='Title')
line('适用对象：首次配置红人工作流看板字幕服务的使用者。本文带你找到指定抓取工具、创建仅供该工具使用的 Token、测试一次字幕，并把 Token 保存到看板。截图中的红色圈框按数字顺序操作。')
line('本次示范使用 scrape-creators / best-youtube-transcripts-scraper。API Token 是密钥，不要发给他人，也不要上传到 GitHub。')
doc.add_heading('第一步 找到字幕抓取工具', 1)
line('登录 Apify，打开 Best Youtube Transcripts Scraper。确认页面标题与下图一致。这里的 Video URLs 可填公开 YouTube 视频链接。')
fig(annotated('apify-guide-actor-input.png', [(245, 13, 633, 58), (292, 263, 986, 313), (676, 394, 829, 440), (243, 453, 331, 492)]), '图 1：确认工具名称；输入视频链接；展开 Run options 设置费用上限；Start 才会运行。')
line('若只想配置 Token，可先不要点 Start。运行字幕抓取可能产生 Apify 费用。')

doc.add_page_break()
doc.add_heading('第二步 创建专用 API Token', 1)
line('在 Apify 左侧点击 Settings，再进入 API & Integrations。在 Personal API tokens 区域点击 Create a new token。')
fig(annotated('apify-guide-token-settings.png', [(8, 667, 200, 698), (451, 111, 595, 145), (1625, 299, 1817, 338), (252, 438, 1810, 510)]), '图 2：依次进入设置、API 与集成、创建新 Token。最下方是本次创建的专用 Token，截图中已遮挡。')
line('建议命名为“红人工作流看板｜YouTube 字幕监测”。权限只选该 Actor 的 Run 和 List runs，并允许访问该次运行的默认存储；不要授予整个账号的全部权限。创建后复制 Token，妥善保存。')
line('如果今后怀疑 Token 泄露，在此页面撤销旧 Token，重新创建并更新看板。')

doc.add_page_break()
doc.add_heading('第三步 测试一条公开视频', 1)
line('回到字幕工具的 Input 页，填一条公开视频。展开 Run options，将 Maximum cost per run 设为 0.05 美元，再点 Start。本步骤会实际调用服务并可能收费；不测试时可跳过。')
fig(annotated('apify-guide-run.png', [(238, 59, 936, 90), (239, 111, 335, 152), (240, 271, 1853, 435)]), '图 3：状态 Succeeded 表示运行成功；Output 显示字幕结果。本次示例实际费用为 0.001 美元。')
line('若出现失败，先查看 Runs 页中的错误提示；也可能是视频没有公开字幕。一次失败不代表 Token 无效。')

doc.add_page_break()
doc.add_heading('第四步 保存到红人工作流看板', 1)
line('登录看板，打开“设置 → YouTube API → YouTube 字幕服务 · Apify”。把 Token 粘贴到 Apify API Token 框，点击“测试账号连接”。连接测试不会启动字幕抓取。成功后开启按需抓取，并点击“保存字幕设置”。')
fig(annotated('apify-guide-app-settings.png', [(28, 609, 221, 648), (289, 375, 618, 418), (298, 474, 1779, 541), (298, 545, 558, 579), (1547, 729, 1659, 767), (1672, 729, 1795, 767)]), '图 4：设置入口、字幕服务、Token 输入框、启用开关、测试连接及保存按钮。图中 Token 已隐藏。')
line('保存后 Token 框显示“已保存，留空不更换”。本次示例每轮最多读取 5 个视频字幕、最多 AI 分析 5 组候选。之后你可以按费用承受能力调整上限。')
doc.add_heading('常见问题', 2)
line('“测试连接”失败：检查复制的 Token、网络，以及是否给这个指定 Actor 授予读取权限。不要为了让测试通过直接授予整个 Apify 账号的全部权限。')
line('“没有字幕”：该工具读取已有公开字幕，不负责把无字幕的音频转写成文字。')
line('费用：看板给单次字幕运行设置 0.05 美元上限，但这不等于整个监测流程所有平台费用的总上限。需要费用控制时，先小范围检测。')

doc.save(OUT)
print(OUT)
